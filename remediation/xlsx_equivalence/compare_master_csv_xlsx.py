#!/usr/bin/env python3
"""Semantic-equivalence check: master cluster/ellipse CSV vs the informational XLSX.

The XLSX at

    remediation/corrected/event_global_max_algorithm/
        scorch_new_algorithm_master_cluster_ellipse_event_global_max.xlsx

was an informational duplicate of the canonical CSV beside it.  This script is
the reproducible evidence behind the ``approved_removal`` record for that path
in ``remediation/corrected_outputs/SHA256_MANIFEST.json``.  It re-derives the
comparison from a *fresh* extraction of the workbook on every run: nothing is
cached and no previous result is consulted.

Run from the repository root::

    python remediation/xlsx_equivalence/compare_master_csv_xlsx.py

The workbook is not required to be present in the repository.  Once the
approved removal has been applied the canonical copy lives in the processed-data
deposit; point the script at an extracted deposit with::

    SCORCH_DATA_DIR=<extracted deposit root> \
        python remediation/xlsx_equivalence/compare_master_csv_xlsx.py

Pass criterion (declared, not discovered)
-----------------------------------------
Equivalence PASSES iff all of the following hold:

* identical shape, column names and column order;
* ALL three declared key columns present (a missing key column is a FAILURE,
  never a silently narrowed key);
* the declared key unique in BOTH the CSV and the XLSX (without uniqueness the
  positional row-order check would not prove row identity);
* identical row keys and row order;
* every non-numeric cell equal exactly;
* every integer-valued column equal exactly;
* identical NaN and Inf masks;
* for every pair of corresponding finite float64 values, the **relative**
  difference ``|a-b| / max(|a|,|b|)`` is <= ``REL_TOL`` (1e-12);
* the workbook is inert with respect to an ENUMERATED list of properties: no
  formulas, external links or external relationship targets, no defined names,
  no VBA archive or macro parts, no charts, images or cell comments, no merged
  cell ranges, no conditional formatting, no data validations, no hyperlinks,
  no worksheet tables, no auto filters, no frozen panes, no pivot caches, and
  exactly one visible sheet.

The inertness claim is scoped to that list, which the run emits as
``inertness_properties_checked``. It is NOT a claim that the workbook carries no
presentation information whatsoever: cell number formats, cell styles, column
widths, print settings and sheet protection are NOT constrained, and the
observed number formats are reported rather than required to be empty. A date
column legitimately carries a display format that a CSV cannot represent, so
demanding zero formats would be a false criterion. None of the unchecked
properties can affect semantic equivalence here, because every cell VALUE is
compared directly and date columns are normalized to ISO strings first.

An *absolute* tolerance is deliberately not used.  The numeric columns span
about nine orders of magnitude, so a fixed absolute bound would demand
sub-float64 precision on the large columns and is not a meaningful criterion.
ULP distances are reported as a diagnostic only: an Excel round-trip is a
decimal serialisation at fixed significant digits, which bounds relative error
rather than ULP distance.

Exit codes: 0 PASS, 1 FAIL, 2 inputs unavailable.
"""
from __future__ import annotations

import json
import os
import sys
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]

CSV_REL = ("remediation/corrected/event_global_max_algorithm/"
           "scorch_new_algorithm_master_cluster_ellipse_event_global_max.csv")
XLSX_REL = ("remediation/corrected/event_global_max_algorithm/"
            "scorch_new_algorithm_master_cluster_ellipse_event_global_max.xlsx")
DEPOSIT_CSV_REL = ("catalogs/"
                   "scorch_new_algorithm_master_cluster_ellipse_event_global_max.csv")
DEPOSIT_XLSX_REL = ("catalogs/"
                    "scorch_new_algorithm_master_cluster_ellipse_event_global_max.xlsx")

SHEET_NAME = "master_cluster_ellipse"
KEY_COLUMNS = ("new_event_id", "date", "cluster_id")
REL_TOL = 1e-12


def _resolve(repo_rel: str, deposit_rel: str) -> Path | None:
    """Repository copy first, then an extracted deposit via SCORCH_DATA_DIR."""
    candidate = REPO / repo_rel
    if candidate.is_file():
        return candidate
    data_dir = os.environ.get("SCORCH_DATA_DIR")
    if data_dir:
        candidate = Path(data_dir).expanduser() / deposit_rel
        if candidate.is_file():
            return candidate
    return None


def workbook_structure(xlsx_path: Path) -> dict:
    """Structural inertness of the workbook, from a fresh open of the file."""
    import openpyxl

    with zipfile.ZipFile(xlsx_path) as zf:
        parts = sorted(zf.namelist())
        rels = ""
        if "xl/_rels/workbook.xml.rels" in parts:
            rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    try:
        formulas = 0
        comments = 0
        images = 0
        charts = 0
        merged = 0
        cond_fmt = 0
        validations = 0
        hyperlinks = 0
        tables = 0
        autofilters = 0
        frozen = 0
        number_formats = set()
        for ws in wb.worksheets:
            images += len(getattr(ws, "_images", []) or [])
            charts += len(getattr(ws, "_charts", []) or [])
            merged += len(getattr(ws, "merged_cells", ()) and
                          ws.merged_cells.ranges or ())
            cond_fmt += sum(1 for _ in (getattr(ws, "conditional_formatting", ())
                                        or ()))
            dv = getattr(ws, "data_validations", None)
            validations += len(getattr(dv, "dataValidation", []) or [])
            tables += len(getattr(ws, "tables", {}) or {})
            af = getattr(ws, "auto_filter", None)
            if af is not None and getattr(af, "ref", None):
                autofilters += 1
            if getattr(ws, "freeze_panes", None):
                frozen += 1
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f" or (
                            isinstance(cell.value, str)
                            and cell.value.startswith("=")):
                        formulas += 1
                    if getattr(cell, "comment", None) is not None:
                        comments += 1
                    if getattr(cell, "hyperlink", None) is not None:
                        hyperlinks += 1
                    number_formats.add(cell.number_format)
        info = {
            "merged_cell_ranges": merged,
            "conditional_formatting_ranges": cond_fmt,
            "data_validations": validations,
            "hyperlinks": hyperlinks,
            "worksheet_tables": tables,
            "auto_filters": autofilters,
            "frozen_panes": frozen,
            "pivot_caches": len(getattr(wb, "_pivots", []) or []),
            # Reported, NOT part of the inertness claim - see the note below.
            "number_formats_observed": sorted(number_formats),
            "zip_parts": parts,
            "sheet_names": list(wb.sheetnames),
            "sheet_count": len(wb.sheetnames),
            "sheet_states": {ws.title: ws.sheet_state for ws in wb.worksheets},
            "hidden_sheets": [ws.title for ws in wb.worksheets
                              if ws.sheet_state != "visible"],
            "formula_cells": formulas,
            "cell_comments": comments,
            "images": images,
            "charts": charts,
            "defined_names": list(wb.defined_names),
            "external_link_parts": [p for p in parts
                                    if p.startswith("xl/externalLink")],
            "external_rel_targets": rels.count('TargetMode="External"'),
            "vba_archive": wb.vba_archive is not None,
            "macro_parts": [p for p in parts if p.endswith(".bin")],
        }
    finally:
        wb.close()

    # The inertness claim is EXACTLY this enumerated list of properties and no
    # more. Cell number formats and cell styles are deliberately NOT in it:
    # a date column legitimately carries a display format, which a CSV cannot
    # represent, so requiring zero number formats would be false. Formats are
    # reported in `number_formats_observed` instead. They cannot affect semantic
    # equivalence here because every cell VALUE is compared directly, and dates
    # are normalized to ISO strings before comparison. Any claim of "no
    # presentation information" must be read as scoped to the checked list.
    info["inertness_properties_checked"] = [
        "exactly one sheet", "no hidden sheets", "no formula cells",
        "no cell comments", "no images", "no charts", "no defined names",
        "no external link parts", "no external relationship targets",
        "no VBA archive", "no macro parts", "no merged cell ranges",
        "no conditional formatting", "no data validations", "no hyperlinks",
        "no worksheet tables", "no auto filters", "no frozen panes",
        "no pivot caches",
    ]
    info["inertness_properties_not_checked"] = [
        "cell number formats (reported, not constrained)",
        "cell styles: fonts, fills, borders, alignment",
        "column widths and row heights",
        "print settings and page setup",
        "sheet protection",
    ]
    info["inert"] = (
        info["sheet_count"] == 1
        and not info["hidden_sheets"]
        and info["formula_cells"] == 0
        and info["cell_comments"] == 0
        and info["images"] == 0
        and info["charts"] == 0
        and not info["defined_names"]
        and not info["external_link_parts"]
        and info["external_rel_targets"] == 0
        and not info["vba_archive"]
        and not info["macro_parts"]
        and info["merged_cell_ranges"] == 0
        and info["conditional_formatting_ranges"] == 0
        and info["data_validations"] == 0
        and info["hyperlinks"] == 0
        and info["worksheet_tables"] == 0
        and info["auto_filters"] == 0
        and info["frozen_panes"] == 0
        and info["pivot_caches"] == 0
    )
    return info


def _ulp_distance(a, b):
    """Distance in representable float64 steps between two finite arrays."""
    import numpy as np

    ai = a.astype(np.float64).view(np.int64)
    bi = b.astype(np.float64).view(np.int64)
    # map sign-magnitude ordering onto a monotone two's-complement ordering
    ai = np.where(ai < 0, np.int64(np.iinfo(np.int64).min) - ai, ai)
    bi = np.where(bi < 0, np.int64(np.iinfo(np.int64).min) - bi, bi)
    return np.abs(ai - bi)


def compare(csv_path: Path, xlsx_path: Path) -> dict:
    import numpy as np
    import pandas as pd

    csv_df = pd.read_csv(csv_path)
    xlsx_df = pd.read_excel(xlsx_path, sheet_name=SHEET_NAME)

    result: dict = {
        "shape_csv": list(csv_df.shape),
        "shape_xlsx": list(xlsx_df.shape),
        "shape_equal": csv_df.shape == xlsx_df.shape,
        "columns_equal": list(csv_df.columns) == list(xlsx_df.columns),
        "csv_only_columns": sorted(set(csv_df.columns) - set(xlsx_df.columns)),
        "xlsx_only_columns": sorted(set(xlsx_df.columns) - set(csv_df.columns)),
        "failures": [],
    }
    if not result["columns_equal"]:
        result["failures"].append("column names or order differ")
        return result
    if not result["shape_equal"]:
        result["failures"].append("shapes differ")
        return result

    # normalise dates to ISO strings so a datetime-typed Excel column and a
    # string-typed CSV column are compared on their content, not their dtype
    for frame in (csv_df, xlsx_df):
        for col in frame.columns:
            if "datetime" in str(frame[col].dtype):
                frame[col] = frame[col].dt.strftime("%Y-%m-%d")

    # The declared key must be REAL and it must be a key. Silently narrowing it
    # to whichever declared columns happen to exist would let the row-identity
    # check degrade to nothing (in the limit, an empty key makes every row
    # identical and the order check vacuously true).
    keys = list(KEY_COLUMNS)
    missing_keys = [c for c in keys if c not in csv_df.columns]
    result["key_columns_declared"] = keys
    result["key_columns_missing"] = missing_keys
    if missing_keys:
        result["failures"].append(f"declared key columns absent: {missing_keys}")
        return result
    result["key_columns"] = keys
    csv_key = csv_df[keys].astype(str).agg("|".join, axis=1)
    xlsx_key = xlsx_df[keys].astype(str).agg("|".join, axis=1)
    result["key_unique_csv"] = bool(csv_key.is_unique)
    result["key_unique_xlsx"] = bool(xlsx_key.is_unique)
    result["key_order_identical"] = bool((csv_key.values == xlsx_key.values).all())
    result["first_key"] = str(csv_key.iloc[0])
    result["last_key"] = str(csv_key.iloc[-1])

    # Uniqueness is a PASS CONDITION, not a diagnostic. A duplicated key means
    # the positional row-order comparison below is not proof of row identity.
    if not result["key_unique_csv"]:
        dupes = sorted(csv_key[csv_key.duplicated()].unique())[:5]
        result["failures"].append(
            f"key {keys} is not unique in the CSV (e.g. {dupes})")
    if not result["key_unique_xlsx"]:
        dupes = sorted(xlsx_key[xlsx_key.duplicated()].unique())[:5]
        result["failures"].append(
            f"key {keys} is not unique in the XLSX (e.g. {dupes})")
    if not (result["key_unique_csv"] and result["key_unique_xlsx"]):
        return result

    # Exact row order, positionally - unchanged.
    if not result["key_order_identical"]:
        result["failures"].append("row keys or row order differ")
        return result

    numeric_cols, string_cols = [], []
    for col in csv_df.columns:
        if (pd.api.types.is_numeric_dtype(csv_df[col])
                and pd.api.types.is_numeric_dtype(xlsx_df[col])):
            numeric_cols.append(col)
        else:
            string_cols.append(col)

    string_diff = {}
    for col in string_cols:
        a = csv_df[col].astype(str).values
        b = xlsx_df[col].astype(str).values
        n = int((a != b).sum())
        if n:
            string_diff[col] = n
    result["string_columns"] = len(string_cols)
    result["string_column_names"] = string_cols
    result["string_differing_cells"] = int(sum(string_diff.values()))
    result["string_differing_columns"] = string_diff
    if string_diff:
        result["failures"].append("non-numeric cells differ")

    nan_mismatch, inf_mismatch = [], []
    integer_cols, float_cols = [], []
    integer_diff = {}
    for col in numeric_cols:
        a = csv_df[col].to_numpy(dtype=np.float64, na_value=np.nan)
        b = xlsx_df[col].to_numpy(dtype=np.float64, na_value=np.nan)
        if not np.array_equal(np.isnan(a), np.isnan(b)):
            nan_mismatch.append(col)
        if not np.array_equal(np.isinf(a), np.isinf(b)):
            inf_mismatch.append(col)
        finite = np.isfinite(a) & np.isfinite(b)
        both_int = (np.all(a[finite] == np.floor(a[finite]))
                    and np.all(b[finite] == np.floor(b[finite])))
        if both_int:
            integer_cols.append(col)
            n = int((a[finite] != b[finite]).sum())
            if n:
                integer_diff[col] = n
        else:
            float_cols.append(col)

    result["numeric_columns"] = len(numeric_cols)
    result["integer_valued_columns"] = len(integer_cols)
    result["float_columns"] = len(float_cols)
    result["nan_mask_identical"] = not nan_mismatch
    result["inf_mask_identical"] = not inf_mismatch
    result["nan_mask_mismatch_columns"] = nan_mismatch
    result["inf_mask_mismatch_columns"] = inf_mismatch
    result["integer_differing_cells"] = int(sum(integer_diff.values()))
    result["integer_differing_columns"] = integer_diff
    if nan_mismatch:
        result["failures"].append("NaN masks differ")
    if inf_mismatch:
        result["failures"].append("Inf masks differ")
    if integer_diff:
        result["failures"].append("integer-valued cells differ")

    total_cells = 0
    differing = 0
    differing_columns = {}
    worst_abs = 0.0
    worst_rel = 0.0
    worst_abs_ctx = None
    worst_rel_ctx = None
    ulps = []
    per_col_max_ulp = {}
    for col in numeric_cols:
        a = csv_df[col].to_numpy(dtype=np.float64, na_value=np.nan)
        b = xlsx_df[col].to_numpy(dtype=np.float64, na_value=np.nan)
        finite = np.isfinite(a) & np.isfinite(b)
        a, b = a[finite], b[finite]
        total_cells += a.size
        neq = a != b
        n = int(neq.sum())
        if not n:
            continue
        differing += n
        differing_columns[col] = n
        da = np.abs(a[neq] - b[neq])
        scale = np.maximum(np.abs(a[neq]), np.abs(b[neq]))
        rel = np.where(scale > 0, da / scale, 0.0)
        u = _ulp_distance(a[neq], b[neq])
        ulps.append(u)
        per_col_max_ulp[col] = int(u.max())
        i = int(np.argmax(da))
        if float(da[i]) > worst_abs:
            worst_abs = float(da[i])
            worst_abs_ctx = {"column": col,
                             "abs": float(da[i]),
                             "rel": float(rel[i]),
                             "magnitude": float(scale[i]),
                             "ulp": int(u[i])}
        j = int(np.argmax(rel))
        if float(rel[j]) > worst_rel:
            worst_rel = float(rel[j])
            worst_rel_ctx = {"column": col,
                             "abs": float(da[j]),
                             "rel": float(rel[j]),
                             "magnitude": float(scale[j]),
                             "ulp": int(u[j])}

    result["numeric_cells_compared"] = int(total_cells)
    result["numeric_differing_cells"] = int(differing)
    result["numeric_differing_columns"] = len(differing_columns)
    result["numeric_differing_column_counts"] = differing_columns
    result["worst_absolute_difference"] = worst_abs
    result["worst_absolute_context"] = worst_abs_ctx
    result["worst_relative_difference"] = worst_rel
    result["worst_relative_context"] = worst_rel_ctx
    result["relative_tolerance"] = REL_TOL
    result["relative_margin_factor"] = (REL_TOL / worst_rel) if worst_rel else None

    if ulps:
        allu = np.concatenate(ulps)
        buckets = {
            "1": int((allu == 1).sum()),
            "2": int((allu == 2).sum()),
            "3-4": int(((allu >= 3) & (allu <= 4)).sum()),
            "5-64": int(((allu >= 5) & (allu <= 64)).sum()),
            ">64": int((allu > 64).sum()),
        }
        result["worst_ulp"] = int(allu.max())
        result["ulp_histogram"] = buckets
        result["worst_ulp_columns"] = sorted(
            (c for c in per_col_max_ulp
             if per_col_max_ulp[c] == max(per_col_max_ulp.values())))
    else:
        result["worst_ulp"] = 0
        result["ulp_histogram"] = {}
        result["worst_ulp_columns"] = []

    if worst_rel > REL_TOL:
        result["failures"].append(
            f"worst relative difference {worst_rel!r} exceeds {REL_TOL!r}")
    return result


def main() -> int:
    csv_path = _resolve(CSV_REL, DEPOSIT_CSV_REL)
    xlsx_path = _resolve(XLSX_REL, DEPOSIT_XLSX_REL)
    if csv_path is None or xlsx_path is None:
        missing = [r for r, p in ((CSV_REL, csv_path), (XLSX_REL, xlsx_path))
                   if p is None]
        print("inputs unavailable: " + ", ".join(missing))
        print("set SCORCH_DATA_DIR to an extracted processed-data deposit "
              "to compare the deposited copies")
        return 2

    import hashlib
    identity = {}
    for label, path in (("csv", csv_path), ("xlsx", xlsx_path)):
        data = path.read_bytes()
        identity[label] = {"bytes": len(data),
                           "sha256": hashlib.sha256(data).hexdigest()}

    structure = workbook_structure(xlsx_path)
    result = compare(csv_path, xlsx_path)
    if not structure["inert"]:
        result["failures"].append("workbook carries logic or linkage")

    passed = not result["failures"]
    report = {
        "csv_source": "repository" if csv_path == REPO / CSV_REL else "deposit",
        "xlsx_source": "repository" if xlsx_path == REPO / XLSX_REL else "deposit",
        "byte_identity_as_read": identity,
        "workbook_structure": structure,
        "comparison": result,
        "verdict": "PASS" if passed else "FAIL",
    }
    json.dump(report, sys.stdout, indent=2, sort_keys=True)
    print()
    print(report["verdict"])
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
